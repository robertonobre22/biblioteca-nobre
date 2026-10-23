from pathlib import Path
from tempfile import TemporaryDirectory
from docx import Document
from converter_docx import convert

def run():
    title="O Homem que Aprendeu a Forjar Silêncio"
    with TemporaryDirectory() as tmp:
        root=Path(tmp)
        doc=Document()
        doc.add_paragraph(title)
        for i in range(1,8):
            doc.add_paragraph(f"Parte {i} — Teste de preservação")
            doc.add_paragraph(f"Parágrafo {i}. Pontuação, travessão — e acentos: í é.")
            doc.add_paragraph("Outra linha com texto   e espaços internos.")
        doc.add_paragraph("Epílogo — Teste")
        doc.add_paragraph("O fim — preservado.")
        source=root/"teste.docx"
        doc.save(source)
        report=convert(source,root/"output",title,"teste-darius")
        assert report["sections"]==8
        assert report["body_paragraphs"]==15
        import json
        book=json.loads((root/"output/livro.json").read_text(encoding="utf-8"))
        assert book["chapters"][0]["paragraphs"][0]=="Parágrafo 1. Pontuação, travessão — e acentos: í é."
        assert book["chapters"][-1]["paragraphs"]==["O fim — preservado."]
        print("Teste de conversão e preservação: OK")
if __name__=="__main__":run()
